import io
import os
import re
import base64
import requests
import zipfile
from datetime import datetime, timedelta
from typing import List, Optional, Dict, Any
from fastapi import FastAPI, HTTPException, Response
from fastapi.responses import HTMLResponse, RedirectResponse
from pydantic import BaseModel, Field
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, borders

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import sys
import secrets
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from edm_service import SAMSAPI_BASE_URL, DEFAULT_SAMSAPI_KEY, SamsApi, resolve_pending, safe_name

app = FastAPI(
    title="Remittance Portal API",
    description="SamsApi 실시간 연동 (미결번호→확정전표(jrn/list) 정확 매칭 V60)",
    version="60.0.0"
)

# 로그인(HTTP Basic). 아이디/비밀번호는 저장소가 공개이므로 코드가 아닌 환경변수로만 설정한다.
# Vercel: Settings → Environment Variables 에 PORTAL_USER / PORTAL_PASSWORD 등록
PORTAL_USER = os.environ.get("PORTAL_USER", "")
PORTAL_PASSWORD = os.environ.get("PORTAL_PASSWORD", "")

@app.middleware("http")
async def require_login(request, call_next):
    if not (PORTAL_USER and PORTAL_PASSWORD):
        # Vercel(외부)에서는 로그인 설정이 없으면 열지 않음. 로컬 실행은 기존처럼 허용
        if os.environ.get("VERCEL"):
            return Response("PORTAL_USER / PORTAL_PASSWORD 환경변수가 설정되지 않아 포털을 열 수 없습니다.", status_code=503, media_type="text/plain; charset=utf-8")
        return await call_next(request)
    auth = request.headers.get("authorization", "")
    if auth.startswith("Basic "):
        try:
            user, _, pw = base64.b64decode(auth[6:]).decode("utf-8").partition(":")
        except Exception:
            user, pw = "", ""
        if secrets.compare_digest(user.encode(), PORTAL_USER.encode()) and secrets.compare_digest(pw.encode(), PORTAL_PASSWORD.encode()):
            return await call_next(request)
    return Response("로그인이 필요합니다.", status_code=401, media_type="text/plain; charset=utf-8",
                    headers={"WWW-Authenticate": 'Basic realm="Heung-A Remittance Portal", charset="UTF-8"'})

manual_payment_dates_db: Dict[str, str] = {}

REGULAR_SUPPLIERS = {
    "한라시스템", "해동구명설비(주)", "한라레벨(주)-한라IMS(주)", "정양엔지니어링",
    "해양금속(주)", "(주)마린테크니컬엔지니어링", "대림엔지니어링(주)", "씨스펙",
    "화진기업(주)", "주식회사 금강항해통신", "(주)해양무선", "주식회사 선양",
    "세연테크", "(주)종합해사청학공장", "(주)이강공사", "(주)매일마린",
    "(주)화일통상", "(주)와이에이치마린", "비아이산업주식회사", "세인베스텍(미우라서비스코리아)",
    "베스트마린코프", "금양엔지니어링", "(주)상봉코포레이션", "(주)디에스케이",
    "주식회사 케이피에스", "(주)해바다", "(주)케이씨", "충무전기공업사", "(주)그린-씨"
}

def get_payment_days_by_amount(amount: float) -> int:
    if amount <= 2_000_000: return 90
    elif amount <= 5_000_000: return 60
    elif amount <= 10_000_000: return 70
    elif amount <= 20_000_000: return 80
    elif amount <= 50_000_000: return 90
    elif amount <= 100_000_000: return 100
    else: return 110

def calculate_payment_date(occur_date_str: str, request_date_str: str, vendor_name: str, amount: float) -> str:
    if not occur_date_str or len(occur_date_str) < 8: occur_date_str = datetime.today().strftime("%Y-%m-%d")
    is_regular = any(supplier in vendor_name for supplier in REGULAR_SUPPLIERS)
    if is_regular: base_target_dt = datetime.strptime(occur_date_str, "%Y-%m-%d") + timedelta(days=get_payment_days_by_amount(amount))
    else:
        if request_date_str and len(request_date_str) >= 8: base_target_dt = datetime.strptime(request_date_str, "%Y-%m-%d")
        else: base_target_dt = datetime.strptime(occur_date_str, "%Y-%m-%d") + timedelta(days=30)
    weekday = base_target_dt.weekday()
    days_to_add_map = {0: 1, 1: 0, 2: 2, 3: 1, 4: 0, 5: 3, 6: 2}
    return (base_target_dt + timedelta(days=days_to_add_map[weekday])).strftime("%Y-%m-%d")

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

class PendingSearchQuery(BaseModel):
    branch_code: str = Field(default="본사")
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    pay_start_date: Optional[str] = None
    pay_end_date: Optional[str] = None
    account_codes: Optional[List[str]] = Field(default_factory=list)
    vendor_code: Optional[str] = None
    pending_no: Optional[str] = None
    unsettled_only: bool = True
    use_mock: Optional[bool] = False
    api_key: Optional[str] = None
    selected_pending_nos: Optional[List[str]] = None
    saved_dates: Optional[Dict[str, str]] = Field(default_factory=dict)

class PaymentDateSaveRequest(BaseModel):
    pending_no: str
    target_payment_date: str

class DebugRequest(BaseModel):
    pending_no: str
    api_key: Optional[str] = None

def get_mock_pending_data(payload: PendingSearchQuery) -> List[dict]:
    items = [
        {"pending_no": "APS202609140021-0001", "account_code": "2002", "account_name": "외상매입금(외화)", "vendor_code": "006717", "vendor_name": "CHINA MARINE SHIPPING AGENCY SHANDONG CO.,LTD WEIFANG BRANCH", "occur_date": "2026-09-14", "acc_date": "2026-09-14", "payment_request_date": "2026-09-14", "currency": "USD", "exchange_rate": 1346.40, "occur_amount": 3763.30, "balance_amount": 3763.30, "krw_balance": 5066903.0, "journal_no": "S202609140041"},
        {"pending_no": "APS202607280020-0001", "account_code": "2002", "account_name": "외상매입금(외화)", "vendor_code": "007438", "vendor_name": "WEIFANG YACHUANG SUPPLY CHAIN CO., LTD", "occur_date": "2026-07-28", "acc_date": "2026-07-28", "payment_request_date": "2026-07-28", "currency": "USD", "exchange_rate": 1548.40, "occur_amount": 1942.94, "balance_amount": 1942.94, "krw_balance": 3008449.0, "journal_no": "S202607280044"}
    ]
    for item in items:
        auto_date = calculate_payment_date(item["occur_date"], item["payment_request_date"], item["vendor_name"], item["krw_balance"])
        item["auto_payment_date"] = auto_date
        item["scheduled_payment_date"] = payload.saved_dates.get(item["pending_no"]) or item["auto_payment_date"]
    return items

def fetch_real_loan_data(payload: PendingSearchQuery) -> List[dict]:
    active_key = payload.api_key.strip() if payload.api_key and payload.api_key.strip() else DEFAULT_SAMSAPI_KEY
    headers = {"X-API-Key": active_key, "Authorization": f"Bearer {active_key}", "Content-Type": "application/json"}
    api_url = f"{SAMSAPI_BASE_URL}/api/v1/loan/currentstatelist"
    target_dt = payload.end_date if (payload.end_date and payload.end_date != "string") else datetime.today().strftime("%Y-%m-%d")
    req_body = {"company_code": "HASL", "target_date": target_dt.replace("-", ""), "language_gubun": "KO"}
    try:
        res = requests.post(api_url, headers=headers, json=req_body, timeout=10)
        if res.status_code == 200 and res.json().get("success"):
            parsed_items = []
            for raw in (res.json().get("data") or []):
                bal = float(raw.get("balance_amount") or 0)
                if bal <= 0: continue
                loan_id = str(raw.get("loand_id") or raw.get("group_settled_number") or "").strip()
                if not loan_id: continue
                from_dt = str(raw.get("from_date") or "").strip()
                from_dt = f"{from_dt[:4]}-{from_dt[4:6]}-{from_dt[6:]}" if len(from_dt)==8 else from_dt
                to_dt = str(raw.get("to_date") or "").strip()
                to_dt = f"{to_dt[:4]}-{to_dt[4:6]}-{to_dt[6:]}" if len(to_dt)==8 else to_dt
                vendor_name = str(raw.get("financial_customer_name") or raw.get("direct_customer_name") or "").strip()
                auto_date = calculate_payment_date(from_dt, to_dt, vendor_name, bal)
                parsed_items.append({
                    "pending_no": loan_id, "account_code": "LOAN", "account_name": "차입금",
                    "vendor_code": raw.get("financial_customer_code", ""), "vendor_name": vendor_name, 
                    "occur_date": from_dt, "acc_date": from_dt, "payment_request_date": to_dt,
                    "currency": raw.get("currency_code", "KRW"), "exchange_rate": 1.0,
                    "occur_amount": float(raw.get("loan_amount") or 0), "balance_amount": bal, "krw_balance": bal,
                    "auto_payment_date": auto_date, "scheduled_payment_date": payload.saved_dates.get(loan_id) or auto_date,
                    "journal_no": "-"
                })
            return parsed_items
        return []
    except: return []

def fetch_real_pending_data(payload: PendingSearchQuery) -> List[dict]:
    active_key = payload.api_key.strip() if payload.api_key and payload.api_key.strip() else DEFAULT_SAMSAPI_KEY
    headers = {"X-API-Key": active_key, "Authorization": f"Bearer {active_key}", "Content-Type": "application/json"}
    api_url = f"{SAMSAPI_BASE_URL}/api/v1/ntstl/list"
    target_dt = payload.end_date if (payload.end_date and payload.end_date != "string") else datetime.today().strftime("%Y-%m-%d")
    acc_code_param = [ac for ac in (payload.account_codes or []) if ac.isdigit()]
    cust_code_param = [payload.vendor_code.strip()] if (payload.vendor_code and payload.vendor_code.strip() and payload.vendor_code.strip() != "string") else []
    
    req_body = {
        "company_code": "HASL",
        "target_date": target_dt.replace("-", ""),
        "type_account_code": acc_code_param,
        "type_customer_code": cust_code_param
    }
    try:
        res = requests.post(api_url, headers=headers, params={"page": 1, "pageSize": 2000}, json=req_body, timeout=10)
        if res.status_code == 200 and res.json().get("success"):
            parsed_items = []
            for raw in (res.json().get("data") or []):
                krw_bal = float(raw.get("local_amount_bal") or raw.get("functional_amount_bal") or 0)
                bal = float(raw.get("occur_amount_bal") or raw.get("local_amount_bal") or 0)
                if bal <= 0 or krw_bal <= 0: continue
                pending_no = str(raw.get("not_settled_number") or "").strip()
                if not pending_no: continue
                occur_dt = str(raw.get("occur_date") or "").strip()
                occur_dt = f"{occur_dt[:4]}-{occur_dt[4:6]}-{occur_dt[6:]}" if len(occur_dt)==8 else occur_dt
                due_dt = str(raw.get("due_date") or "").strip()
                due_dt = f"{due_dt[:4]}-{due_dt[4:6]}-{due_dt[6:]}" if len(due_dt)==8 else due_dt
                vendor_name = str(raw.get("customer_name") or "").strip()
                auto_date = calculate_payment_date(occur_dt, due_dt, vendor_name, krw_bal)
                parsed_items.append({
                    "pending_no": pending_no, "account_code": raw.get("account_code", ""), "account_name": raw.get("account_name", ""),
                    "vendor_code": raw.get("customer_code", ""), "vendor_name": vendor_name, "occur_date": occur_dt, "acc_date": occur_dt,
                    "payment_request_date": due_dt, "currency": raw.get("currency_code", "KRW"), "exchange_rate": float(raw.get("occur_exchange_rate") or 1.0),
                    "occur_amount": float(raw.get("occur_amount_ocr") or 0), "balance_amount": bal, "krw_balance": krw_bal,
                    "auto_payment_date": auto_date, "scheduled_payment_date": payload.saved_dates.get(pending_no) or auto_date,
                    "journal_no": "-"
                })
            return parsed_items
        return [{"error_msg": f"SAMSAPI 데이터 수신 실패: {res.text}"}]
    except Exception as e: return [{"error_msg": f"SAMSAPI 네트워크 연결 오류: {str(e)}"}]

# 전표번호는 조회 시 매칭하지 않음 (확정전표 전체 대조는 수 분 소요).
# 선택 건만 edm_service.resolve_pending 으로 정확히 매칭한다 → /api/pending/resolve-journals, ZIP 다운로드
def fetch_combined_dataset(payload: PendingSearchQuery) -> List[dict]:
    selected = [a.strip() for a in payload.account_codes] if payload.account_codes else []
    has_loan = ("차입금" in selected) or ("LOAN" in selected) or (len(selected) == 0)
    has_pending = any(a in selected for a in ["2001", "2002", "미지급금"]) or (len(selected) == 0) or (has_loan and len(selected) > 1)

    if payload.use_mock: return get_mock_pending_data(payload)
    
    combined = []
    if has_pending: combined.extend(fetch_real_pending_data(payload))
    if has_loan: combined.extend(fetch_real_loan_data(payload))
    return combined

def filter_data(payload: PendingSearchQuery, data: List[dict]) -> List[dict]:
    if data and data[0].get("error_msg"): return data
    if payload.start_date and payload.start_date.strip() not in ["", "string"]: data = [item for item in data if item["occur_date"] >= payload.start_date]
    if payload.end_date and payload.end_date.strip() not in ["", "string"]: data = [item for item in data if item["occur_date"] <= payload.end_date]
    if payload.pay_start_date and payload.pay_start_date.strip() not in ["", "string"]: data = [item for item in data if item["scheduled_payment_date"] >= payload.pay_start_date]
    if payload.pay_end_date and payload.pay_end_date.strip() not in ["", "string"]: data = [item for item in data if item["scheduled_payment_date"] <= payload.pay_end_date]
        
    if payload.account_codes and len(payload.account_codes) > 0:
        filtered_by_acc = []
        for item in data:
            item_code = str(item.get("account_code", "")).lower()
            item_name = str(item.get("account_name", "")).lower()
            match = False
            for target in payload.account_codes:
                t = target.strip().lower()
                if t == "차입금" or t == "loan":
                    if "loan" in item_code or "차입" in item_name or "loan" in item_name: match = True; break
                elif t in item_code or t in item_name: match = True; break
            if match: filtered_by_acc.append(item)
        data = filtered_by_acc

    if payload.vendor_code and payload.vendor_code.strip() not in ["", "string"]:
        v_code = payload.vendor_code.strip().lower()
        data = [item for item in data if v_code in item.get("vendor_code", "").lower() or v_code in item.get("vendor_name", "").lower()]
    if payload.pending_no and payload.pending_no.strip() not in ["", "string"]:
        p_no = payload.pending_no.strip().lower()
        data = [item for item in data if p_no in item.get("pending_no", "").lower()]
    return data

@app.get("/portal", response_class=HTMLResponse)
def render_portal_ui():
    return f"""
    <!DOCTYPE html>
    <html lang="ko">
    <head>
        <meta charset="UTF-8">
        <title>흥아해운 송금 관리 포털</title>
        <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        <style>
            body {{ background-color: #f4f6f9; font-family: 'Malgun Gothic', sans-serif; }}
            .navbar {{ background-color: #1a365d; }}
            .card {{ border-radius: 8px; border: none; box-shadow: 0 2px 8px rgba(0,0,0,0.06); }}
            .table-responsive {{ overflow-x: auto; }}
            .resizable-table {{ table-layout: fixed; width: 100%; min-width: 1300px; border-collapse: collapse; }}
            .resizable-table th {{ position: relative; background-color: #2b4c7e; color: white; padding: 10px; border: 1px solid #dee2e6; user-select: none; }}
            .resizer {{ width: 6px; height: 100%; position: absolute; right: 0; top: 0; cursor: col-resize; z-index: 1; }}
            .resizer:hover, .resizer.resizing {{ background-color: #ffc107; border-right: 2px solid #e0a800; }}
            .btn-excel {{ background-color: #1d6f42; color: white; }}
            .btn-remit {{ background-color: #c0392b; color: white; }}
            .btn-zip {{ background-color: #6f42c1; color: white; }}
            .btn-mock {{ background-color: #6c757d; color: white; border-color: #6c757d; }}
            .btn-scan {{ background-color: #0d6efd; color: white; }}
            .bg-summary {{ background-color: #fffbeb; }}
            .date-input {{ background-color: #e8f5e9; border: 1px solid #4caf50; color: #1b5e20; font-weight: bold; width: 100%; }}
            .chk-group {{ background-color: #f8f9fa; padding: 8px 15px; border-radius: 6px; border: 1px solid #dee2e6; box-shadow: inset 0 1px 3px rgba(0,0,0,0.05); }}
        </style>
    </head>
    <body class="p-3">
        <nav class="navbar navbar-dark px-4 py-3 rounded mb-4 d-flex justify-content-between">
            <span class="navbar-brand mb-0 h1 fw-bold">🚢 흥아해운 미결 포털 <span class="badge bg-primary fs-6 ms-2">미결→전표 정확 매칭(V60) 🟢</span></span>
        </nav>
        
        <div class="card p-3 mb-4">
            <h5 class="fw-bold text-secondary mb-3">🔍 상세 미결 & 차입금 조회 조건 (회사코드: HASL)</h5>
            <div class="row g-3 mb-2">
                <div class="col-md-2"><label class="form-label text-secondary fw-bold">발생 시작일</label><input type="date" class="form-control" id="startDate" value="2025-01-01"></div>
                <div class="col-md-2"><label class="form-label text-secondary fw-bold">발생 종료일</label><input type="date" class="form-control" id="endDate" value=""></div>
                <div class="col-md-2"><label class="form-label fw-bold text-success">지불예정 시작일</label><input type="date" class="form-control border-success" id="payStartDate" value=""></div>
                <div class="col-md-2"><label class="form-label fw-bold text-success">지불예정 종료일</label><input type="date" class="form-control border-success" id="payEndDate" value=""></div>
            </div>

            <div class="row g-3 mb-2 align-items-end">
                <div class="col-md-6">
                    <label class="form-label text-secondary fw-bold">계정과목 선택 (스위치 - 다중 선택 가능)</label>
                    <div class="chk-group d-flex gap-4 align-items-center">
                        <div class="form-check form-switch"><input class="form-check-input acc-chk" type="checkbox" value="2001" id="acc2001"><label class="form-check-label fw-bold" for="acc2001">2001 (원화)</label></div>
                        <div class="form-check form-switch"><input class="form-check-input acc-chk" type="checkbox" value="2002" id="acc2002" checked><label class="form-check-label fw-bold" for="acc2002">2002 (외화)</label></div>
                        <div class="form-check form-switch"><input class="form-check-input acc-chk" type="checkbox" value="미지급금" id="accUnpaid" checked><label class="form-check-label fw-bold" for="accUnpaid">미지급금</label></div>
                        <div class="form-check form-switch"><input class="form-check-input acc-chk" type="checkbox" value="차입금" id="accLoan"><label class="form-check-label fw-bold text-primary" for="accLoan">🏦 차입금</label></div>
                    </div>
                </div>

                <div class="col-md-3"><label class="form-label text-secondary fw-bold">거래처/금융기관 (코드/명)</label><input type="text" class="form-control" id="vendorCode" placeholder="예: CHINA"></div>
                <div class="col-md-3"><label class="form-label text-secondary fw-bold">미결/차입 번호</label><input type="text" class="form-control" id="pendingNo" placeholder="예: APS2026..."></div>
            </div>

            <div class="d-flex justify-content-end gap-2 mt-3">
                <button class="btn btn-mock fw-bold px-4" onclick="loadPendingData(true)">MOCK조회(테스트)</button>
                <button class="btn btn-primary fw-bold px-5" onclick="loadPendingData(false)">조회(API)</button>
                <button class="btn btn-excel fw-bold px-4" onclick="downloadExcel()">리스트(엑셀)</button>
                <button class="btn btn-remit fw-bold px-4" onclick="downloadRemittanceForm()">📄 BNK 외화송금신청서(워드)</button>
                <button class="btn btn-scan fw-bold px-4" id="btnResolve" onclick="resolveJournals()">선택 전표번호 조회</button>
                <button class="btn btn-zip fw-bold px-4" onclick="downloadEdmZip()">선택항목 증빙 ZIP</button>
            </div>
        </div>

        <div class="card p-3 mb-4">
            <div class="d-flex flex-wrap align-items-center gap-2 mb-3 p-2 rounded" style="background-color: #e8f5e9;">
                <span class="fw-bold text-success">📅 지불/만기예정일 일괄 변경 (체크한 건)</span>
                <input type="date" class="form-control form-control-sm" id="bulkDate" style="width: 160px;">
                <button class="btn btn-sm btn-success fw-bold" onclick="applyBulkDate()">선택 건에 일괄 적용</button>
                <button class="btn btn-sm btn-outline-secondary fw-bold" onclick="resetBulkDate()">선택 건 자동산정일로 되돌리기</button>
            </div>
            <div class="table-responsive">
                <table class="table table-hover align-middle border text-center resizable-table" style="font-size: 0.88rem;" id="pendingTable">
                    <thead>
                        <tr>
                            <th style="width: 45px;"><input type="checkbox" id="selectAll" onclick="toggleAll(this)"></th>
                            <th style="width: 160px;">미결/차입번호</th>
                            <th style="width: 140px;">전표번호</th>
                            <th style="width: 130px;">계정/차입구분</th>
                            <th style="width: 180px;">거래처/금융기관</th>
                            <th style="width: 70px;">통화</th>
                            <th style="width: 120px;">외화(원화잔액)</th>
                            <th style="width: 140px;">원화환산액</th>
                            <th style="width: 100px;">자동산정일</th>
                            <th style="width: 140px; background-color: #1b5e20;">지불/만기예정일</th>
                            <th style="width: 100px;">작업</th>
                        </tr>
                    </thead>
                    <tbody id="pendingTableBody"><tr><td colspan="11" class="py-4 text-muted">조회 버튼을 눌러 데이터를 불러오세요.</td></tr></tbody>
                </table>
            </div>
        </div>

        <div class="card p-3 border-warning">
            <h5 class="fw-bold text-warning mb-3">📊 지불 / 만기 계획 요약 (통화별 합계)</h5>
            <div class="table-responsive">
                <table class="table table-bordered text-center align-middle">
                    <thead class="bg-light"><tr><th>통화 (Currency)</th><th>통화별 원화 환산 합계 (KRW Converted)</th></tr></thead>
                    <tbody id="summaryTableBody"></tbody>
                    <tfoot><tr class="bg-summary fw-bold fs-5 text-danger"><td class="text-end pe-4">총 원화 환산 지불 계획 금액 :</td><td id="grandTotalKrw">0 원</td></tr></tfoot>
                </table>
            </div>
        </div>

        <!-- 엑스레이 팝업 모달 -->
        <div class="modal fade" id="debugModal" tabindex="-1" aria-hidden="true">
          <div class="modal-dialog modal-xl">
            <div class="modal-content">
              <div class="modal-header bg-dark text-white">
                <h5 class="modal-title fw-bold">🔍 SAMSAPI 원본 데이터 엑스레이 결과</h5>
                <button type="button" class="btn-close btn-close-white" data-bs-dismiss="modal" aria-label="Close"></button>
              </div>
              <div class="modal-body bg-light">
                <textarea id="debugResultText" class="form-control" style="height: 500px; font-family: monospace; font-size: 0.85rem;" readonly></textarea>
              </div>
            </div>
          </div>
        </div>

        <script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
        <script>
            function initResizableTable() {{
                const table = document.getElementById('pendingTable');
                const ths = table.querySelectorAll('th');
                ths.forEach(th => {{
                    if(th.querySelector('.resizer')) return; 
                    const resizer = document.createElement('div'); resizer.classList.add('resizer'); th.appendChild(resizer);
                    let startX, startWidth;
                    resizer.addEventListener('mousedown', function(e) {{
                        startX = e.pageX; startWidth = th.offsetWidth; resizer.classList.add('resizing');
                        function mouseMoveHandler(e) {{ const newWidth = startWidth + (e.pageX - startX); if (newWidth > 40) {{ th.style.width = newWidth + 'px'; th.style.minWidth = newWidth + 'px'; }} }}
                        function mouseUpHandler() {{ resizer.classList.remove('resizing'); document.removeEventListener('mousemove', mouseMoveHandler); document.removeEventListener('mouseup', mouseUpHandler); }}
                        document.addEventListener('mousemove', mouseMoveHandler); document.addEventListener('mouseup', mouseUpHandler); e.stopPropagation();
                    }});
                }});
            }}

            function buildSearchPayload(useMock = false) {{
                const savedDatesObj = JSON.parse(localStorage.getItem('manualPaymentDates') || '{{}}');
                const selectedAccs = [];
                document.querySelectorAll('.acc-chk:checked').forEach(chk => selectedAccs.push(chk.value));

                return {{
                    branch_code: "본사",
                    start_date: document.getElementById("startDate").value,
                    end_date: document.getElementById("endDate").value,
                    pay_start_date: document.getElementById("payStartDate").value,
                    pay_end_date: document.getElementById("payEndDate").value,
                    account_codes: selectedAccs,
                    vendor_code: document.getElementById("vendorCode").value,
                    pending_no: document.getElementById("pendingNo").value,
                    unsettled_only: true,
                    use_mock: useMock,
                    saved_dates: savedDatesObj
                }};
            }}

            function toggleAll(source) {{ document.querySelectorAll('.row-chk').forEach(chk => chk.checked = source.checked); }}

            async function loadPendingData(useMock = false) {{
                const payload = buildSearchPayload(useMock);
                const tbody = document.getElementById("pendingTableBody");
                const summaryBody = document.getElementById("summaryTableBody");
                
                tbody.innerHTML = '<tr><td colspan="11" class="py-4 text-primary fw-bold">데이터를 조회 중입니다...</td></tr>';
                summaryBody.innerHTML = ""; document.getElementById("grandTotalKrw").innerText = "0 원";
                
                try {{
                    const res = await fetch('/api/pending/search-and-schedule', {{ method: 'POST', headers: {{'Content-Type': 'application/json'}}, body: JSON.stringify(payload) }});
                    const text = await res.text();
                    let data = JSON.parse(text);
                    
                    if(data.length > 0 && data[0].error_msg) {{ tbody.innerHTML = `<tr><td colspan="11" class="py-4 text-danger fw-bold fs-5">${{data[0].error_msg}}</td></tr>`; return; }}
                    if(data.length === 0) {{ tbody.innerHTML = '<tr><td colspan="11" class="py-4 text-muted fw-bold">검색 조건에 일치하는 데이터가 없습니다.</td></tr>'; return; }}

                    let summary = {{}}; let grandTotalKrw = 0; tbody.innerHTML = "";
                    data.forEach(item => {{
                        let curr = item.currency;
                        if(!summary[curr]) summary[curr] = {{ krw: 0 }};
                        summary[curr].krw += item.krw_balance;
                        grandTotalKrw += item.krw_balance;

                        const tr = document.createElement("tr");
                        tr.innerHTML = `
                            <td><input type="checkbox" class="row-chk form-check-input" value="${{item.pending_no}}"></td>
                            <td class="text-primary fw-bold text-break">${{item.pending_no}}</td>
                            <td class="text-dark fw-bold text-break" id="jrn-${{item.pending_no}}">${{item.journal_no || '-'}}</td>
                            <td class="text-break"><span class="badge bg-light text-dark border">${{item.account_name || item.account_code}}</span></td>
                            <td class="fw-bold text-break">${{item.vendor_name}}</td>
                            <td><span class="badge ${{curr === 'KRW' ? 'bg-secondary' : 'bg-danger'}}">${{curr}}</span></td>
                            <td class="text-end pe-2 text-break">${{Number(item.balance_amount).toLocaleString()}}</td>
                            <td class="fw-bold text-end pe-2 text-break">${{Number(item.krw_balance).toLocaleString()}} 원</td>
                            <td><span class="text-muted">${{item.auto_payment_date}}</span></td>
                            <td><input type="date" class="form-control form-control-sm text-center date-input" id="date-${{item.pending_no}}" data-auto="${{item.auto_payment_date}}" value="${{item.scheduled_payment_date}}"></td>
                            <td>
                                <button class="btn btn-sm btn-success fw-bold w-100 mb-1" onclick="saveDate('${{item.pending_no}}')">저장</button>
                                <button class="btn btn-sm btn-scan fw-bold w-100" onclick="debugRow('${{item.pending_no}}')">🔍엑스레이</button>
                            </td>
                        `;
                        tbody.appendChild(tr);
                    }});
                    for(const [curr, amounts] of Object.entries(summary)) {{
                        const tr = document.createElement("tr");
                        tr.innerHTML = `<td class="fw-bold text-primary" colspan="2">${{curr}}</td><td class="text-end pe-4 fw-bold" colspan="9">${{Number(amounts.krw).toLocaleString()}} 원</td>`;
                        summaryBody.appendChild(tr);
                    }}
                    document.getElementById("grandTotalKrw").innerText = Number(grandTotalKrw).toLocaleString() + " 원";
                }} catch(e) {{ tbody.innerHTML = `<tr><td colspan="11" class="py-4 text-danger fw-bold">🚨 오류 발생: ${{e.message}}</td></tr>`; }}
            }}
            
            // 체크한 건의 지불예정일을 한 번에 변경하고 브라우저에 저장 (다음 조회 때도 유지)
            function setDatesForChecked(getDate) {{
                const checkedBoxes = document.querySelectorAll('.row-chk:checked');
                if (checkedBoxes.length === 0) {{ alert("변경할 건을 선택해주세요."); return 0; }}
                const savedObj = JSON.parse(localStorage.getItem('manualPaymentDates') || '{{}}');
                checkedBoxes.forEach(cb => {{
                    const input = document.getElementById(`date-${{cb.value}}`);
                    const newDate = getDate(input);
                    if (!input || !newDate) return;
                    input.value = newDate;
                    if (newDate === input.dataset.auto) delete savedObj[cb.value]; else savedObj[cb.value] = newDate;
                }});
                localStorage.setItem('manualPaymentDates', JSON.stringify(savedObj));
                return checkedBoxes.length;
            }}

            function applyBulkDate() {{
                const newDate = document.getElementById("bulkDate").value;
                if (!newDate) {{ alert("적용할 날짜를 먼저 선택해주세요."); return; }}
                const n = setDatesForChecked(() => newDate);
                if (n) alert(`${{n}}건의 지불/만기예정일을 ${{newDate}}로 변경했습니다.`);
            }}

            function resetBulkDate() {{
                const n = setDatesForChecked(input => input.dataset.auto);
                if (n) alert(`${{n}}건을 자동산정일로 되돌렸습니다.`);
            }}

            async function saveDate(pendingNo) {{
                const newDate = document.getElementById(`date-${{pendingNo}}`).value;
                const savedObj = JSON.parse(localStorage.getItem('manualPaymentDates') || '{{}}');
                savedObj[pendingNo] = newDate;
                localStorage.setItem('manualPaymentDates', JSON.stringify(savedObj));

                const res = await fetch('/api/pending/save-payment-date', {{ method: 'POST', headers: {{'Content-Type': 'application/json'}}, body: JSON.stringify({{ pending_no: pendingNo, target_payment_date: newDate }}) }});
                alert((await res.json()).message + "\\n(수정된 지불예정일이 정상 고정되었습니다.)");
            }}

            async function debugRow(pendingNo) {{
                const modal = new bootstrap.Modal(document.getElementById('debugModal'));
                const textArea = document.getElementById('debugResultText');
                textArea.value = "서버에서 API 원본 데이터를 스캔 중입니다...";
                modal.show();
                
                try {{
                    const res = await fetch('/api/pending/debug-raw', {{
                        method: 'POST', headers: {{'Content-Type': 'application/json'}},
                        body: JSON.stringify({{ pending_no: pendingNo }})
                    }});
                    const data = await res.json();
                    textArea.value = JSON.stringify(data, null, 2);
                }} catch(e) {{
                    textArea.value = "오류 발생: " + e.message;
                }}
            }}

            function downloadExcel() {{
                fetch('/api/pending/export-plan-excel', {{ method: 'POST', headers: {{'Content-Type': 'application/json'}}, body: JSON.stringify(buildSearchPayload(false)) }})
                .then(res => res.blob()).then(blob => {{ const a = document.createElement('a'); a.href = window.URL.createObjectURL(blob); a.download = `전체_지불_계획리스트.xlsx`; a.click(); }});
            }}

            function downloadRemittanceForm() {{
                const checkedBoxes = document.querySelectorAll('.row-chk:checked');
                if (checkedBoxes.length === 0) {{ alert("BNK 외화송금신청서를 출력할 건을 선택해주세요."); return; }}
                const payload = buildSearchPayload(false);
                payload.selected_pending_nos = Array.from(checkedBoxes).map(cb => cb.value);
                fetch('/api/pending/export-remittance-form', {{ method: 'POST', headers: {{'Content-Type': 'application/json'}}, body: JSON.stringify(payload) }})
                .then(res => res.blob()).then(blob => {{ const a = document.createElement('a'); a.href = window.URL.createObjectURL(blob); a.download = `BNK_부산은행_외화송금신청서.docx`; a.click(); }});
            }}
            
            async function resolveJournals() {{
                const checkedBoxes = document.querySelectorAll('.row-chk:checked');
                if (checkedBoxes.length === 0) {{ alert("전표번호를 조회할 건을 선택해주세요."); return; }}
                const payload = buildSearchPayload(false);
                payload.selected_pending_nos = Array.from(checkedBoxes).map(cb => cb.value);
                const btn = document.getElementById('btnResolve');
                const originalText = btn.innerText;
                btn.innerText = `전표번호 조회 중... (${{checkedBoxes.length}}건)`; btn.disabled = true;
                try {{
                    const res = await fetch('/api/pending/resolve-journals', {{ method: 'POST', headers: {{'Content-Type': 'application/json'}}, body: JSON.stringify(payload) }});
                    for (const r of await res.json()) {{
                        const cell = document.getElementById(`jrn-${{r.pending_no}}`);
                        if (!cell) continue;
                        cell.innerHTML = r.journal_no
                            ? `${{r.journal_no}}<br><small class="text-muted">증빙 ${{r.edm_count}}개</small>`
                            : `<small class="text-danger">${{r.error || '전표 없음'}}</small>`;
                    }}
                }} catch(e) {{ alert("전표번호 조회 중 오류: " + e.message); }}
                btn.innerText = originalText; btn.disabled = false;
            }}

            function downloadEdmZip() {{
                const checkedBoxes = document.querySelectorAll('.row-chk:checked');
                if (checkedBoxes.length === 0) {{ alert("증빙 자료를 다운로드할 건을 선택해주세요."); return; }}
                const payload = buildSearchPayload(false);
                payload.selected_pending_nos = Array.from(checkedBoxes).map(cb => cb.value);

                const btn = document.querySelector('.btn-zip');
                const originalText = btn.innerText;
                btn.innerText = "전표번호 매칭 + 증빙 다운로드 중..."; btn.disabled = true;

                fetch('/api/pending/export-edm-zip', {{ method: 'POST', headers: {{'Content-Type': 'application/json'}}, body: JSON.stringify(payload) }})
                .then(res => res.blob()).then(blob => {{ 
                    const a = document.createElement('a'); a.href = window.URL.createObjectURL(blob); a.download = `실제_전표_EDM증빙자료.zip`; a.click(); 
                    btn.innerText = originalText; btn.disabled = false;
                }}).catch(() => {{ 
                    alert("다운로드 중 오류가 발생했습니다."); btn.innerText = originalText; btn.disabled = false;
                }});
            }}
            window.onload = function() {{
                // 발생 종료일 기본값 = 오늘 (사용자 PC 기준 날짜, YYYY-MM-DD)
                document.getElementById("endDate").value = new Date().toLocaleDateString('sv-SE');
                initResizableTable(); loadPendingData(false);
            }};
        </script>
    </body>
    </html>
    """

@app.post("/api/pending/debug-raw")
def debug_raw(payload: DebugRequest):
    api = SamsApi(payload.api_key)
    p_no = payload.pending_no.strip()
    result = {}
    try:
        info = api.get_pending_info([p_no]).get(p_no)
        result["1_ntstlinfo_미결상세조회"] = info or "미결번호 없음"
        if info:
            jrn = api.find_journal(p_no, str(info.get("occur_date") or ""), str(info.get("customer_code") or ""))
            result["2_jrn/list_미결번호_일치_확정전표"] = jrn or "일치하는 확정전표 없음"
            if jrn:
                result["3_edm/list_증빙목록"] = api.list_edm(str(jrn.get("journal_number") or "").strip())
    except Exception as e:
        result["오류"] = str(e)
    return result

@app.post("/api/pending/resolve-journals")
def resolve_journals(payload: PendingSearchQuery):
    items = resolve_pending(SamsApi(payload.api_key), payload.selected_pending_nos or [])
    return [{"pending_no": i["pending_no"], "journal_no": i.get("journal_number"), "edm_count": len(i.get("edm_files") or []), "error": i.get("error")} for i in items]

@app.get("/")
def read_root(): return RedirectResponse("/portal")

@app.post("/api/pending/search-and-schedule")
def search_and_schedule_pending(payload: PendingSearchQuery):
    try: return filter_data(payload, fetch_combined_dataset(payload))
    except Exception as e: return [{"error_msg": f"백엔드 처리 오류: {str(e)}"}]

@app.post("/api/pending/save-payment-date")
def save_payment_date(payload: PaymentDateSaveRequest): 
    manual_payment_dates_db[payload.pending_no] = payload.target_payment_date
    return {"status": "success", "message": f"[{payload.pending_no}] 지불예정일이 {payload.target_payment_date}로 정상 저장되었습니다."}

@app.post("/api/pending/export-plan-excel")
def export_plan_excel(payload: PendingSearchQuery):
    filtered_data = filter_data(payload, fetch_combined_dataset(payload))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "전체_리스트"
    ws.append(["지불/만기예정일", "미결/차입번호", "전표번호", "계정/차입구분", "거래처/금융기관", "통화", "환율", "발생/차입금액", "원화환산액", "자동산정일"])
    for cell in ws[1]: cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid"); cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
    for row in filtered_data:
        if row.get("error_msg"): continue
        ws.append([row["scheduled_payment_date"], row["pending_no"], row.get("journal_no", ""), row.get("account_name", ""), row["vendor_name"], row["currency"], row["exchange_rate"], row["balance_amount"], row["krw_balance"], row["auto_payment_date"]])
    stream = io.BytesIO(); wb.save(stream); stream.seek(0)
    return Response(content=stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=List.xlsx"})

@app.post("/api/pending/export-remittance-form")
def export_remittance_form(payload: PendingSearchQuery):
    filtered_data = filter_data(payload, fetch_combined_dataset(payload))
    if payload.selected_pending_nos:
        filtered_data = [item for item in filtered_data if item["pending_no"] in payload.selected_pending_nos]

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    for idx, item in enumerate(filtered_data):
        if idx > 0: doc.add_page_break()

        header_table = doc.add_table(rows=1, cols=2)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        cell_l = header_table.cell(0, 0)
        p_l = cell_l.paragraphs[0]
        run_l1 = p_l.add_run("외화송금신청서\n")
        run_l1.font.name = '맑은 고딕'; run_l1.font.size = Pt(16); run_l1.font.bold = True
        run_l2 = p_l.add_run("APPLICATION FOR REMITTANCE")
        run_l2.font.name = '맑은 고딕'; run_l2.font.size = Pt(10); run_l2.font.bold = True

        cell_r = header_table.cell(0, 1)
        p_r = cell_r.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_r = p_r.add_run("BNK 부산은행  [은행용]")
        run_r.font.name = '맑은 고딕'; run_r.font.size = Pt(13); run_r.font.bold = True
        run_r.font.color.rgb = RGBColor(192, 57, 43)

        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run("지급신청서 및 거래외국환은행 지정확인(신청)서 겸용 Request for designation of correspondent foreign exchange bank and foreign exchange payment")
        run_sub.font.name = '맑은 고딕'; run_sub.font.size = Pt(7.5)

        table = doc.add_table(rows=16, cols=6)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        col_widths = [Inches(0.5), Inches(1.4), Inches(1.3), Inches(2.2), Inches(1.2), Inches(1.2)]
        for row in table.rows:
            for i, w in enumerate(col_widths): row.cells[i].width = w

        table.cell(0, 0).merge(table.cell(2, 0))
        table.cell(0, 0).text = "신\n청\n인"
        table.cell(0, 1).text = "성명 (법인명)\nApplicant"
        table.cell(0, 2).text = "한글: 흥아해운(주)"
        table.cell(0, 3).merge(table.cell(0, 5))
        table.cell(0, 3).text = "영문: Heung-A Shipping Co.,Ltd."

        table.cell(1, 1).merge(table.cell(1, 2))
        table.cell(1, 1).text = "사업자등록번호"
        table.cell(1, 3).merge(table.cell(1, 5))
        table.cell(1, 3).text = "120-81-62522"

        table.cell(2, 1).merge(table.cell(2, 2))
        table.cell(2, 1).text = "주소"
        table.cell(2, 3).merge(table.cell(2, 5))
        table.cell(2, 3).text = "서울특별시 중구 청계천로 8, 5층 (프리미어플레이스)"

        table.cell(3, 0).merge(table.cell(15, 0))
        table.cell(3, 0).text = "신\n청\n내\n용"

        table.cell(3, 1).merge(table.cell(3, 2))
        table.cell(3, 1).text = "약식송금신청"
        table.cell(3, 3).merge(table.cell(3, 5))
        table.cell(3, 3).text = ""

        table.cell(4, 1).merge(table.cell(4, 2))
        table.cell(4, 1).text = "송금방법 (Type)"
        table.cell(4, 3).merge(table.cell(4, 5))
        table.cell(4, 3).text = "■ 해외송금(일반,약식)  □ 중국원화(CNY)송금  □ 타행환송금  □ 국내외화송금"

        table.cell(5, 1).merge(table.cell(5, 2))
        table.cell(5, 1).text = "송금신청액(Amount)"
        table.cell(5, 3).text = f"통화: {item.get('currency', 'USD')}"
        table.cell(5, 4).text = f"금액: {item.get('balance_amount', 0):,.2f}"
        table.cell(5, 5).text = "미화상당액:"

        table.cell(6, 1).merge(table.cell(8, 1))
        table.cell(6, 1).text = "수취인\n(Beneficiary)"
        table.cell(6, 2).text = "성명(업체명)"
        table.cell(6, 3).text = item.get("vendor_name", "")
        table.cell(6, 4).text = "신청인과의 관계"
        table.cell(6, 5).text = ""

        table.cell(7, 2).text = "주소"
        table.cell(7, 3).merge(table.cell(7, 5))
        table.cell(7, 3).text = ""

        table.cell(8, 2).text = "국적"
        table.cell(8, 3).text = "SINGAPORE" if "SINGAPORE" in item.get("vendor_name", "").upper() else ""
        table.cell(8, 4).text = "중국CNY신분증"
        table.cell(8, 5).text = ""

        table.cell(9, 1).merge(table.cell(12, 1))
        table.cell(9, 1).text = "수취거래은행\n(Beneficiary's Bank)"
        table.cell(9, 2).text = "SWIFT BIC"
        table.cell(9, 3).text = "OCBCSGSG" if "UNITED" in item.get("vendor_name", "").upper() or "MARINA" in item.get("vendor_name", "").upper() else ""
        table.cell(9, 4).text = "은행코드"
        table.cell(9, 5).text = ""

        table.cell(10, 2).text = "계좌번호"
        table.cell(10, 3).merge(table.cell(10, 5))
        table.cell(10, 3).text = "503344509301" if "UNITED" in item.get("vendor_name", "").upper() or "MARINA" in item.get("vendor_name", "").upper() else ""

        table.cell(11, 2).text = "은행명"
        table.cell(11, 3).merge(table.cell(11, 5))
        table.cell(11, 3).text = "OCBC BANK, SINGAPORE" if "UNITED" in item.get("vendor_name", "").upper() or "MARINA" in item.get("vendor_name", "").upper() else ""

        table.cell(12, 2).text = "은행주소"
        table.cell(12, 3).merge(table.cell(12, 5))
        table.cell(12, 3).text = ""

        table.cell(13, 1).merge(table.cell(13, 2))
        table.cell(13, 1).text = "송금사유"
        table.cell(13, 3).text = f"{datetime.today().strftime('%y%m%d')}_송금({item.get('pending_no', '')})"
        table.cell(13, 4).text = "수수료부담 (Charges)"
        table.cell(13, 5).text = "각자부담 ■"

        table.cell(14, 1).merge(table.cell(14, 2))
        table.cell(14, 1).text = "수입대금의 경우"
        table.cell(14, 3).text = "H.S. code:"
        table.cell(14, 4).text = "L/C or 계약서 NO:"
        table.cell(14, 5).text = "대응수입예정일:"

        table.cell(15, 1).merge(table.cell(15, 2))
        table.cell(15, 1).text = "중간경유은행"
        table.cell(15, 3).text = ""
        table.cell(15, 4).text = "기타통지사항"
        table.cell(15, 5).text = ""

        for row in table.rows:
            for cell in row.cells:
                set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = '맑은 고딕'; r.font.size = Pt(8.5)

        p_b1 = doc.add_paragraph()
        run_b1 = p_b1.add_run("■ 본인은 귀행 영업점에 비치된 「외환거래기본약관」 및 「전자금융외화송금거래약관」을 열람하고 그 내용에 따를 것을 확약하며 신청합니다.")
        run_b1.font.name = '맑은 고딕'; run_b1.font.size = Pt(8); run_b1.font.bold = True

        p_b2 = doc.add_paragraph()
        run_b2 = p_b2.add_run("■ 본 거래는 미국, UN, EU등이 정한 경제제재 대상자 또는 국가와 관련이 없음을 확인합니다.\n")
        run_b2.font.name = '맑은 고딕'; run_b2.font.size = Pt(8)

        p_sig = doc.add_paragraph()
        p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_sig = p_sig.add_run("예금주명: 흥아해운(주) (인)")
        run_sig.font.name = '맑은 고딕'; run_sig.font.size = Pt(11); run_sig.font.bold = True

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=BNK_BUSAN_BANK_REMITTANCE_APPLICATION.docx"}
    )

@app.post("/api/pending/export-edm-zip")
def export_edm_zip(payload: PendingSearchQuery):
    api = SamsApi(payload.api_key)
    zip_buffer = io.BytesIO()

    with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
        for counter, item in enumerate(resolve_pending(api, payload.selected_pending_nos or []), 1):
            p_no = item["pending_no"]
            folder = safe_name(f"{counter}_[{p_no}]_{item.get('vendor_name') or '알수없음'}"[:80])
            journal_no = item.get("journal_number")
            saved = 0
            for f in item.get("edm_files") or []:
                name = safe_name(f"{f.get('order', '')}_{f.get('filename') or 'file.pdf'}")
                try:
                    zip_file.writestr(f"{folder}/{name}", api.download(f["downloadurl"]))
                    saved += 1
                except Exception as e:
                    zip_file.writestr(f"{folder}/{name}_다운로드실패.txt", str(e).encode("utf-8"))
            if saved == 0:
                if item.get("error"):
                    msg = f"미결번호 [{p_no}] : {item['error']}"
                else:
                    msg = f"미결번호 [{p_no}]의 전표번호({journal_no})를 찾았으나, 해당 전표에 등록된 EDM 증빙 파일이 없습니다."
                zip_file.writestr(f"{folder}_증빙없음.txt", msg.encode("utf-8"))

    zip_buffer.seek(0)
    return Response(content=zip_buffer.getvalue(), media_type="application/zip", headers={"Content-Disposition": "attachment; filename=Real_EDM_Documents.zip"})